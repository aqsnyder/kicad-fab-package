import argparse
import collections
import io
import json
import os
import sys
import tempfile
import types
import unittest
import zipfile
from contextlib import redirect_stdout
from unittest import mock

from kicad_fab_package import export_fab as ef

from ._util import rmtree

GERBER = '''G04 #@! TF.GenerationSoftware,KiCad,Pcbnew,10.0.7*
G04 Created by KiCad (PCBNEW 10.0.7) date 2026-01-01 00:00:00*
%FSLAX46Y46*%
G04 Gerber Fmt 4.6, Leading zero omitted, Abs format (unit mm)*
%MOMM*%
%LPD*%
G01*
G04 APERTURE LIST*
%ADD10C,0.100000*%
%ADD11R,1.000000X2.000000*%
G04 APERTURE END LIST*
D10*
X1000000Y-2000000D02*
X3000000Y-2000000D01*
D11*
X5000000Y-5000000D03*
G36*
X0Y0D02*
X1000000Y0D01*
X1000000Y1000000D01*
X0Y0D01*
G37*
%LPC*%
X6000000Y-6000000D03*
M02*
'''

EXCELLON = '''M48
; DRILL file {KiCad 10.0.7} date 2026-01-01T00:00:00
; FORMAT={-:-/ absolute / metric / decimal}
; #@! TF.FileFunction,Plated,1,4,PTH
FMAT,2
METRIC
; #@! TA.AperFunction,Plated,PTH,ViaDrill
T1C0.300
; #@! TA.AperFunction,Plated,PTH,ComponentDrill
T2C1.000
%
G90
G05
T1
X10.0Y-5.0
X12.0Y-5.0
T2
X20.0Y-8.0
M30
'''

BOARD_VIAS = '''(kicad_pcb
\t(version 20250000)
\t(setup
\t\t(pad_to_mask_clearance 0)
\t\t(capping yes)
\t)
\t(via
\t\t(at 10 5)
\t\t(size 0.6)
\t\t(drill 0.3)
\t\t(layers "F.Cu" "B.Cu")
\t\t(filling yes)
\t\t(net "GND")
\t\t(uuid "v1")
\t)
\t(via blind
\t\t(at 12 5)
\t\t(size 0.6)
\t\t(drill 0.3)
\t\t(layers "F.Cu" "In1.Cu")
\t\t(capping no)
\t\t(net "GND")
\t\t(uuid "v2")
\t)
)
'''


# kicad-cli pcb export step --board-only --user-origin of a board whose Edge.Cuts is one gr_circle, 18.5 mm across
# (KiCad 10.0.7; the product and presentation entities left out). Its CARTESIAN_POINTs reach only y 0..9.25: the
# circle's centre and its seam vertex.
CIRCLE_STEP = '''ISO-10303-21;
DATA;
#11 = AXIS2_PLACEMENT_3D('',#12,#13,#14);
#12 = CARTESIAN_POINT('',(0.,0.,0.));
#13 = DIRECTION('',(0.,0.,1.));
#14 = DIRECTION('',(1.,0.,-0.));
#19 = ( GEOMETRIC_REPRESENTATION_CONTEXT(3)
GLOBAL_UNCERTAINTY_ASSIGNED_CONTEXT((#23)) GLOBAL_UNIT_ASSIGNED_CONTEXT(
(#20,#21,#22)) REPRESENTATION_CONTEXT('Context #1',
  '3D Context with UNIT and UNCERTAINTY') );
#33 = MANIFOLD_SOLID_BREP('',#34);
#34 = CLOSED_SHELL('',(#35,#68,#77));
#35 = ADVANCED_FACE('',(#36),#63,.T.);
#36 = FACE_BOUND('',#37,.F.);
#37 = EDGE_LOOP('',(#38,#48,#55,#56));
#38 = ORIENTED_EDGE('',*,*,#39,.T.);
#39 = EDGE_CURVE('',#40,#42,#44,.T.);
#40 = VERTEX_POINT('',#41);
#41 = CARTESIAN_POINT('',(18.5,9.25,0.));
#42 = VERTEX_POINT('',#43);
#43 = CARTESIAN_POINT('',(18.5,9.25,1.51));
#44 = LINE('',#45,#46);
#45 = CARTESIAN_POINT('',(18.5,9.25,0.));
#46 = VECTOR('',#47,1.);
#47 = DIRECTION('',(0.,0.,1.));
#48 = ORIENTED_EDGE('',*,*,#49,.T.);
#49 = EDGE_CURVE('',#42,#42,#50,.T.);
#50 = CIRCLE('',#51,9.25);
#51 = AXIS2_PLACEMENT_3D('',#52,#53,#54);
#52 = CARTESIAN_POINT('',(9.25,9.25,1.51));
#53 = DIRECTION('',(0.,0.,1.));
#54 = DIRECTION('',(1.,0.,-0.));
#55 = ORIENTED_EDGE('',*,*,#39,.F.);
#56 = ORIENTED_EDGE('',*,*,#57,.F.);
#57 = EDGE_CURVE('',#40,#40,#58,.T.);
#58 = CIRCLE('',#59,9.25);
#59 = AXIS2_PLACEMENT_3D('',#60,#61,#62);
#60 = CARTESIAN_POINT('',(9.25,9.25,0.));
#61 = DIRECTION('',(0.,0.,1.));
#62 = DIRECTION('',(1.,0.,-0.));
#63 = CYLINDRICAL_SURFACE('',#64,9.25);
#64 = AXIS2_PLACEMENT_3D('',#65,#66,#67);
#65 = CARTESIAN_POINT('',(9.25,9.25,0.));
#66 = DIRECTION('',(-0.,-0.,-1.));
#67 = DIRECTION('',(1.,0.,-0.));
#72 = PLANE('',#73);
#73 = AXIS2_PLACEMENT_3D('',#74,#75,#76);
#74 = CARTESIAN_POINT('',(0.,0.,0.));
#75 = DIRECTION('',(0.,0.,1.));
#76 = DIRECTION('',(1.,0.,-0.));
ENDSEC;
END-ISO-10303-21;
'''

# one arc edge of a board whose right side bulges 2 mm out of a 20 mm chord (radius 26 mm, centre 24 mm inside the
# board), as kicad-cli writes it: the circle turns about -Z, so the short arc from (130, -100) to (130, -120) passes
# x = 132. {edge} is filled in with the edge's vertices and sense.
ARC_STEP = '''DATA;
#49 = EDGE_CURVE('',{edge});
#42 = VERTEX_POINT('',#43);
#43 = CARTESIAN_POINT('',(130.,-100.,1.51));
#50 = VERTEX_POINT('',#51);
#51 = CARTESIAN_POINT('',(130.,-120.,1.51));
#52 = CIRCLE('',#53,26.);
#53 = AXIS2_PLACEMENT_3D('',#54,#55,#56);
#54 = CARTESIAN_POINT('',(106.,-110.,1.51));
#55 = DIRECTION('',(0.,0.,-1.));
#56 = DIRECTION('',(0.923076923077,0.384615384615,0.));
ENDSEC;
'''

# a cubic B-spline, a rational one written as a complex entity, a circle behind a SURFACE_CURVE, and (malformed, to
# show 2-D points are left out) an edge whose curve is a p-curve's 2-D spline
CURVES_STEP = '''DATA;
#1 = EDGE_CURVE('',#2,#4,#6,.T.);
#2 = VERTEX_POINT('',#3);
#3 = CARTESIAN_POINT('',(0.,0.,0.));
#4 = VERTEX_POINT('',#5);
#5 = CARTESIAN_POINT('',(10.,0.,0.));
#6 = B_SPLINE_CURVE_WITH_KNOTS('',3,(#3,#7,#8,#5),.UNSPECIFIED.,.F.,.F.,(4,4),(0.,1.),
  .PIECEWISE_BEZIER_KNOTS.);
#7 = CARTESIAN_POINT('',(3.,4.,0.));
#8 = CARTESIAN_POINT('',(7.,4.,0.));
#20 = EDGE_CURVE('',#2,#4,#21,.T.);
#21 = ( BOUNDED_CURVE() B_SPLINE_CURVE(2,(#3,#22,#5),.UNSPECIFIED.,.F.,.F.)
B_SPLINE_CURVE_WITH_KNOTS((3,3),(0.,1.),.PIECEWISE_BEZIER_KNOTS.) CURVE() GEOMETRIC_REPRESENTATION_ITEM()
RATIONAL_B_SPLINE_CURVE((1.,0.7,1.)) REPRESENTATION_ITEM('') );
#22 = CARTESIAN_POINT('',(5.,-6.,0.));
#30 = EDGE_CURVE('',#31,#31,#33,.T.);
#31 = VERTEX_POINT('',#32);
#32 = CARTESIAN_POINT('',(25.,0.,0.));
#33 = SURFACE_CURVE('',#34,(#40),.PCURVE_S1.);
#34 = CIRCLE('',#35,5.);
#35 = AXIS2_PLACEMENT_3D('',#36,#37,$);
#36 = CARTESIAN_POINT('',(20.,0.,0.));
#37 = DIRECTION('',(0.,0.,1.));
#40 = PCURVE('',#41,#42);
#42 = DEFINITIONAL_REPRESENTATION('',(#43),#44);
#43 = B_SPLINE_CURVE_WITH_KNOTS('',1,(#45,#46),.UNSPECIFIED.,.F.,.F.,(2,2),(0.,1.),.UNSPECIFIED.);
#45 = CARTESIAN_POINT('',(-50.,0.));
#46 = CARTESIAN_POINT('',(50.,99.));
#50 = EDGE_CURVE('',#2,#4,#51,.T.);
#51 = SEAM_CURVE('',#43,(#40),.PCURVE_S1.);
ENDSEC;
'''


def extents(points):
    """(xmin, xmax, ymin, ymax, zmin, zmax), rounded to 1 nm."""
    return tuple(round(f(p[k] for p in points), 6) + 0.0 for k in range(3) for f in (min, max))


def write(path, text, newline='\n'):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline=newline) as f:
        f.write(text)
    return path


class TmpDirTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='kfp_test_')

    def tearDown(self):
        rmtree(self.tmp)


class XyrsRulesTest(unittest.TestCase):
    def test_rotation_top_is_kicad_orientation(self):
        for ori, want in ((0, 0), (90, 90), (180, 180), (270, 270), (-90, 270), (450, 90)):
            self.assertEqual(ef.xyrs_rotation(ori, 1)[1], want, ori)

    def test_rotation_bottom_adds_180(self):
        for ori, want in ((0, 180), (90, 270), (180, 0), (270, 90), (-90, 90)):
            self.assertEqual(ef.xyrs_rotation(ori, 2)[1], want, ori)

    def test_rotation_rounding(self):
        raw, rot = ef.xyrs_rotation(89.996, 1)
        self.assertAlmostEqual(raw, 89.996)
        self.assertEqual(rot, 90)
        self.assertEqual(ef.xyrs_rotation(359.7, 1)[1], 0)   # rounds to 360 -> 0

    def test_point_origin_bottom_left_y_up(self):
        origin = (100.0, 50.0)        # (left, bottom) in KiCad mm, Y down
        self.assertEqual(ef.xyrs_point(100.0, 50.0, origin), (0.0, 0.0))
        x, y = ef.xyrs_point(100.0 + 25.4, 50.0 - 25.4, origin)   # 1 inch right, 1 inch up
        self.assertAlmostEqual(x, 1000.0)
        self.assertAlmostEqual(y, 1000.0)

    def test_non_part_and_dnp_words(self):
        self.assertTrue(ef.NON_PART_RE.match('FID1'))
        self.assertTrue(ef.NON_PART_RE.match('H12'))
        self.assertFalse(ef.NON_PART_RE.match('HS1'))
        self.assertTrue(ef.DNP_WORD_RE.search('ABC1001dnpR'))
        self.assertFalse(ef.DNP_WORD_RE.search('DN1'))

    def test_non_part_rule_is_configurable(self):
        import re
        self.assertTrue(ef.is_non_part('FID1', 'Custom:Fid'))
        self.assertTrue(ef.is_non_part('P1', 'Fiducial:Fiducial_1mm'))
        # designator only (no library match), library only (no designator match), neither
        self.assertFalse(ef.is_non_part('P1', 'Fiducial:Fiducial_1mm', ef.NON_PART_RE, ()))
        self.assertFalse(ef.is_non_part('FID1', 'Custom:Fid', None, ef.NON_PART_LIBS))
        self.assertTrue(ef.is_non_part('MH2', 'Mech:Hole', re.compile(r'MH\d+$'), ()))
        self.assertFalse(ef.is_non_part('FID1', 'Fiducial:F', None, ()))

    def test_fiducials_line_by_designator_or_kicad_library(self):
        self.assertTrue(ef.is_fiducial('FID3', 'Custom:Mark'))
        self.assertTrue(ef.is_fiducial('FD1', 'Fiducial:Fiducial_1mm_Mask2mm'))
        self.assertFalse(ef.is_fiducial('H1', 'MountingHole:MountingHole_3.2mm_M3'))
        self.assertFalse(ef.is_fiducial('FD1', 'Custom:Fiducial_1mm'))       # library name, not footprint name


class FilledViaChecksTest(unittest.TestCase):
    def test_board_without_filled_vias(self):
        groups = {('none', 'none'): {(1.0, 2.0)}}
        self.assertTrue(ef.filled_vias_seen(set(), groups))
        self.assertTrue(ef.filled_vias_found([], []))
        self.assertFalse(ef.filled_vias_seen(set(), groups, required=True))
        self.assertFalse(ef.filled_vias_seen(set(), {}, required=True))
        self.assertFalse(ef.filled_vias_found([], [], required=True))

    def test_board_with_filled_vias(self):
        want = {(1.0, 2.0)}
        groups = {('filled', 'capped'): {(1.0, 2.0)}, ('none', 'none'): {(3.0, 4.0)}}
        for required in (False, True):
            self.assertTrue(ef.filled_vias_seen(want, groups, required))
            self.assertFalse(ef.filled_vias_seen(want, {('none', 'none'): want | {(5.0, 5.0)}}, required))
            self.assertTrue(ef.filled_vias_found([], [{'x': 1.0}], required))
            self.assertFalse(ef.filled_vias_found(['(1.0, 2.0) 0 matches'], [{'x': 1.0}], required))


class NamesTest(unittest.TestCase):
    def test_upload_base_has_no_dots(self):
        self.assertEqual(ef.upload_base('myboard_v1.2.3'), 'myboard_v1_2_3')
        self.assertEqual(ef.upload_base('plain'), 'plain')

    def test_fab_layer_stem(self):
        self.assertEqual(ef.fab_layer_stem('myboard-L3 (+Sig2).gbr', 'myboard'), 'L3_Sig2')
        self.assertEqual(ef.fab_layer_stem('myboard-F_Cu.gbr', 'myboard'), 'F_Cu')
        self.assertEqual(ef.fab_layer_stem('myboard-Edge_Cuts.gm1', 'myboard'), 'Edge_Cuts')
        self.assertEqual(ef.fab_layer_stem('other-B.Mask.gbr', 'myboard'), 'other_B_Mask')

    def test_copper_extensions(self):
        self.assertEqual(ef.copper_upload_ext(1, 'Top', 'Positive'), 'GTL')
        self.assertEqual(ef.copper_upload_ext(6, 'Bot', 'Positive'), 'GBL')
        self.assertEqual(ef.copper_upload_ext(3, 'Inr', 'Positive'), 'G3L')
        self.assertEqual(ef.copper_upload_ext(4, 'Inr', 'Negative'), 'G4P')

    def test_natural_key(self):
        self.assertEqual(sorted(['C10', 'C2', 'C1', 'R1'], key=ef.natural_key), ['C1', 'C2', 'C10', 'R1'])

    def test_non_copper_layer_table(self):
        exts = sorted(v[1] for v in ef.NON_COPPER.values())
        self.assertEqual(exts, ['BOR', 'GBO', 'GBP', 'GBS', 'GTO', 'GTP', 'GTS'])


class GerberTest(TmpDirTest):
    def test_gerber_objects(self):
        g = ef.gerber_objects(write(os.path.join(self.tmp, 'a.gbr'), GERBER))
        self.assertEqual(g['apertures'][10], ('C', [0.1]))
        self.assertEqual(g['apertures'][11], ('R', [1.0, 2.0]))
        self.assertEqual(g['draws'], [(1.0, -2.0, 3.0, -2.0, 10, 'D')])
        self.assertEqual(g['flashes'], [(5.0, -5.0, 11, 'D'), (6.0, -6.0, 11, 'C')])
        self.assertEqual(len(g['regions']), 1)
        pts, pol = g['regions'][0]
        self.assertEqual(pol, 'D')
        self.assertEqual(pts[0], (0.0, 0.0))
        self.assertIn((1.0, 1.0), pts)

    def test_gerber_ops_extents_and_arcs(self):
        flashes, pts = ef.gerber_ops(write(os.path.join(self.tmp, 'a.gbr'), GERBER))
        self.assertEqual(flashes, [(5.0, -5.0), (6.0, -6.0)])
        self.assertIn((3.0, -2.0), pts)
        arc = ('%FSLAX46Y46*%\n%MOMM*%\nG01*\nX1000000Y0D02*\n'
               'G03X-1000000Y0I-1000000J0D01*\nG02X1000000Y0I1000000J0D01*\nM02*\n')
        _, pts = ef.gerber_ops(write(os.path.join(self.tmp, 'arc.gbr'), arc))
        ys = [p[1] for p in pts]
        self.assertAlmostEqual(max(ys), 1.0)    # CCW half circle through (0, 1)
        # the CW arc from (-1, 0) back to (1, 0) about (0, 0) passes (0, 1) as well: no point below y = 0
        self.assertGreaterEqual(min(ys), -1e-9)
        arc_cw = arc.replace('G03X-1000000Y0I-1000000J0D01*\nG02X1000000Y0I1000000J0D01*',
                             'G02X-1000000Y0I-1000000J0D01*')
        _, pts = ef.gerber_ops(write(os.path.join(self.tmp, 'arc_cw.gbr'), arc_cw))
        self.assertAlmostEqual(min(p[1] for p in pts), -1.0)   # CW half circle through (0, -1)

    def test_inch_units(self):
        g = '%FSLAX46Y46*%\n%MOIN*%\nD10*\nX1000000Y1000000D03*\nM02*\n'
        flashes, _ = ef.gerber_ops(write(os.path.join(self.tmp, 'in.gbr'), g))
        self.assertAlmostEqual(flashes[0][0], 25.4)

    def test_x2_attribute_lines(self):
        p = write(os.path.join(self.tmp, 'x2.gbr'), '%TF.FileFunction,Copper,L1,Top*%\n%TA.AperFunction,SMDPad*%\n'
                                                    'G04 comment*\n%TD*%\n%MOMM*%\n')
        self.assertEqual(ef.x2_attribute_lines(p), 3)
        self.assertEqual(ef.x2_attribute_lines(write(os.path.join(self.tmp, 'a.gbr'), GERBER)), 0)

    def test_date_line_filter(self):
        lines = GERBER.splitlines()
        kept = [l for l in lines if not ef.DATE_LINE_RE.search(l)]
        self.assertEqual(len(lines) - len(kept), 1)


class ExcellonTest(TmpDirTest):
    def test_parse(self):
        ex = ef.Excellon(write(os.path.join(self.tmp, 'a.drl'), EXCELLON))
        self.assertEqual(ex.pre[0], 'M48')
        self.assertIn('METRIC', ex.pre)
        self.assertEqual(list(ex.tools), [1, 2])
        self.assertEqual(ex.tools[1]['dia'], 0.3)
        self.assertEqual(ex.tools[1]['lines'][0], '; #@! TA.AperFunction,Plated,PTH,ViaDrill')
        self.assertEqual([(t, xy) for t, xy, _ in ex.hits()], [(1, (10.0, -5.0)), (1, (12.0, -5.0)), (2, (20.0, -8.0))])
        self.assertEqual(ex.body_pre, ['G90', 'G05'])
        self.assertEqual(ex.tail, ['M30'])
        self.assertEqual(ex.slot_lines, [])
        self.assertEqual(ex.eol, '\n')

    def test_split_round_trip_keeps_crlf(self):
        src = write(os.path.join(self.tmp, 'a.drl'), EXCELLON, newline='\r\n')
        ex = ef.Excellon(src)
        self.assertEqual(ex.eol, '\r\n')
        keep = {t: [ln for _, ln in d['hits']] for t, d in ex.tools.items()}
        moved = {1: [keep[1].pop(0)]}
        a, b = os.path.join(self.tmp, 'keep.drl'), os.path.join(self.tmp, 'moved.drl')
        ex.write(a, keep, 'note A')
        ex.write(b, moved, 'note B')
        A, B = ef.Excellon(a), ef.Excellon(b)
        self.assertEqual([xy for _, xy, _ in A.hits()], [(12.0, -5.0), (20.0, -8.0)])
        self.assertEqual([xy for _, xy, _ in B.hits()], [(10.0, -5.0)])
        self.assertEqual(list(B.tools.keys()), [1])          # unused tools are dropped from the header
        with open(b, 'rb') as f:
            raw = f.read()
        self.assertTrue(raw.startswith(b'M48\r\n; note B\r\n'))
        self.assertTrue(raw.endswith(b'M30\r\n'))

    def test_slots_and_errors(self):
        slot = EXCELLON.replace('X20.0Y-8.0', 'G00X20.0Y-8.0\nM15\nG01X21.0Y-8.0\nM16')
        ex = ef.Excellon(write(os.path.join(self.tmp, 's.drl'), slot))
        self.assertTrue(ex.slot_lines)
        with self.assertRaises(ef.StepError):
            ef.Excellon(write(os.path.join(self.tmp, 'bad.drl'), EXCELLON.replace('%\n', '')))
        with self.assertRaises(ef.StepError):
            ef.Excellon(write(os.path.join(self.tmp, 'bad2.drl'), EXCELLON.replace('T1\nX10', 'X10')))


class BoardTextTest(unittest.TestCase):
    def test_board_vias_text_defaults_from_setup(self):
        vias = ef.board_vias_text(BOARD_VIAS)
        self.assertEqual(len(vias), 2)
        v1, v2 = vias
        self.assertEqual((v1['x'], v1['y'], v1['drill']), (10.0, 5.0, 0.3))
        self.assertTrue(v1['filled'])            # explicit
        self.assertTrue(v1['capped'])            # from (setup (capping yes))
        self.assertEqual(v1['kind'], '')
        self.assertEqual(v1['layers'], '"F.Cu" "B.Cu"')
        self.assertFalse(v2['filled'])           # no token, setup has no (filling yes)
        self.assertFalse(v2['capped'])           # explicit no beats the setup default
        self.assertEqual(v2['kind'], 'blind')


class StepExtentsTest(unittest.TestCase):
    def test_round_outline_spans_the_whole_circle(self):
        # the old extents (every CARTESIAN_POINT) saw y 0..9.25 on this board: 'STEP: origin = XYRS corner' failed
        naive = [p for p in ef.step_edge_points(CIRCLE_STEP.replace('EDGE_CURVE', 'EDGE_CURV'))]
        self.assertEqual(extents(naive)[2:4], (0.0, 9.25))
        self.assertEqual(extents(ef.step_edge_points(CIRCLE_STEP)), (0.0, 18.5, 0.0, 18.5, 0.0, 1.51))

    def test_arc_extremes_follow_the_edge_and_its_sense(self):
        short = (130.0, 132.0, -120.0, -100.0, 1.51, 1.51)
        for edge in ('#42,#50,#52,.T.', '#50,#42,#52,.F.'):         # the same short arc, either way round
            self.assertEqual(extents(ef.step_edge_points(ARC_STEP.format(edge=edge))), short, edge)
        # the other arc between the same two points is the rest of the circle: out to x = 80 and y = -136 .. -84
        self.assertEqual(extents(ef.step_edge_points(ARC_STEP.format(edge='#50,#42,#52,.T.'))),
                         (80.0, 130.0, -136.0, -84.0, 1.51, 1.51))

    def test_splines_complex_entities_and_surface_curves(self):
        pts = ef.step_edge_points(CURVES_STEP)
        self.assertEqual(extents(pts), (0.0, 25.0, -6.0, 5.0, 0.0, 0.0))
        self.assertIn((3.0, 4.0, 0.0), pts)                    # a B-spline's control points bound it
        self.assertIn((5.0, -6.0, 0.0), pts)                   # ... also in a complex (rational) entity
        self.assertIn((15.0, 0.0, 0.0), [tuple(round(c, 6) + 0.0 for c in p) for p in pts])   # circle, no ref dir

    def test_no_edges_falls_back_to_every_point(self):
        text = "DATA;\n#1 = CARTESIAN_POINT('',(1.,2.,3.));\n#2 = CARTESIAN_POINT('',(-1.,5.,0.));\nENDSEC;\n"
        self.assertEqual(extents(ef.step_edge_points(text)), (-1.0, 1.0, 2.0, 5.0, 0.0, 3.0))

    def test_step_args(self):
        self.assertEqual(ef.step_args("'',(1.,2.),#5,.T."), ["''", '(1.,2.)', '#5', '.T.'])
        self.assertEqual(ef.step_args("'a, (b)',$"), ["'a, (b)'", '$'])


class PasteTest(unittest.TestCase):
    def test_paste_findings(self):
        fps = {
            'C1': {'pad_box_mm': [0, 0, 1, 1], 'paste_pads': [], 'paste_graphics': []},
            'C2': {'pad_box_mm': [5, 0, 6, 1], 'paste_pads': [], 'paste_graphics': []},
            'R1': {'pad_box_mm': [2, 0, 3, 1], 'paste_pads': [['1', 'F.Paste']], 'paste_graphics': []},
            'U1': {'pad_box_mm': [0, 0, 4, 4], 'paste_pads': [],
                   'paste_graphics': [{'class': 'PCB_SHAPE', 'layer': 'B.Paste', 'uuid': 'u-1',
                                       'bbox_mm': [1, 1, 3, 3]}]},
            'J1': {'pad_box_mm': [8, 0, 9, 1], 'paste_pads': [], 'paste_graphics': []},
        }
        rows = [{'ref': 'C1', 'side': 1, 'pop': 1, 'type': 1}, {'ref': 'C2', 'side': 1, 'pop': 1, 'type': 1},
                {'ref': 'R1', 'side': 1, 'pop': 0, 'type': 1}, {'ref': 'U1', 'side': 2, 'pop': 0, 'type': 1},
                {'ref': 'J1', 'side': 1, 'pop': 1, 'type': 2}]
        paste = {'F': {'flashes': [(0.5, -0.5, 10, 'D'), (2.5, -0.5, 10, 'D'), (5.5, -0.5, 10, 'C')],
                       'regions': [], 'draws': []},
                 'B': {'flashes': [], 'regions': [([(1, -1), (3, -1), (3, -3)], 'D')], 'draws': []}}
        source, plotted, unpasted = ef.paste_findings(fps, rows, paste)
        self.assertEqual(list(source), ['R1', 'U1'])
        self.assertIn('1 pad on a paste layer (1 F.Paste)', source['R1'][0])
        self.assertIn('uuid u-1', source['U1'][0])
        self.assertEqual(dict(plotted['R1']), {'flash': 1})
        self.assertEqual(dict(plotted['U1']), {'region': 1})
        self.assertEqual(unpasted, ['C2'])        # its only flash is clear polarity; J1 is through-hole


def lz_compress_b64(s):
    """LZString.compressToBase64, ported for the round-trip test of export_fab.lz_decompress_b64. A Python port of
    lz-string's compressor (lz-string: MIT License, Copyright (c) 2013 pieroxy,
    https://github.com/pieroxy/lz-string)."""
    alphabet = 'ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789+/='
    bits_per_char = 6
    st = {'val': 0, 'pos': 0}
    data = []

    def put(value, nbits):
        for _ in range(nbits):
            st['val'] = (st['val'] << 1) | (value & 1)
            if st['pos'] == bits_per_char - 1:
                st['pos'] = 0
                data.append(alphabet[st['val']])
                st['val'] = 0
            else:
                st['pos'] += 1
            value >>= 1

    dictionary, to_create = {}, set()
    w, enlarge, dict_size, num_bits = '', 2, 3, 2

    def emit_w():
        nonlocal enlarge, num_bits
        if w in to_create:
            if ord(w[0]) < 256:
                put(0, num_bits)
                put(ord(w[0]), 8)
            else:
                put(1, num_bits)
                put(ord(w[0]), 16)
            enlarge -= 1
            if enlarge == 0:
                enlarge, num_bits = 1 << num_bits, num_bits + 1
            to_create.discard(w)
        else:
            put(dictionary[w], num_bits)
        enlarge -= 1
        if enlarge == 0:
            enlarge, num_bits = 1 << num_bits, num_bits + 1

    for c in s:
        if c not in dictionary:
            dictionary[c] = dict_size
            dict_size += 1
            to_create.add(c)
        wc = w + c
        if wc in dictionary:
            w = wc
        else:
            emit_w()
            dictionary[wc] = dict_size
            dict_size += 1
            w = c
    if w:
        emit_w()
    put(2, num_bits)
    while True:
        st['val'] <<= 1
        if st['pos'] == bits_per_char - 1:
            data.append(alphabet[st['val']])
            break
        st['pos'] += 1
    res = ''.join(data)
    return res + '=' * ((4 - len(res) % 4) % 4)


class LzStringTest(unittest.TestCase):
    def test_round_trip(self):
        samples = ['a', 'hello hello hello', 'abababababababababab', json.dumps({'bom': {'F': [[['C1', 1]]]},
                                                                               'metadata': {'revision': '1.0'}}),
                   'unicode µF Ω 中文', ''.join(chr(32 + (i * 7) % 90) for i in range(3000))]
        for s in samples:
            self.assertEqual(ef.lz_decompress_b64(lz_compress_b64(s)), s, s[:20])

    def test_empty(self):
        self.assertEqual(ef.lz_decompress_b64(''), '')


class CtxTest(TmpDirTest):
    def test_short_tags_paths_with_forward_slashes(self):
        C = ef.Ctx(argparse.Namespace(ibom_pkg=None))
        C.work = os.path.join(self.tmp, 'work')
        C.out = os.path.join(self.tmp, 'out')
        C.top = os.path.join(self.tmp, 'repo')
        s = C.short('%s and %s and %s' % (os.path.join(C.work, 'a', 'b.txt'), os.path.join(C.out, 'fab'), C.top))
        self.assertEqual(s, '<work>/a/b.txt and <out>/fab and <repo>')
        self.assertEqual(C.short(ef.__file__), '<export_fab.py>')
        p = C.portable({'k': [os.path.join(C.out, 'x')]})
        self.assertEqual(p['k'], ['<out>/x'])

    def test_short_repo_wins_when_it_is_the_main_checkout(self):
        C = ef.Ctx(argparse.Namespace(ibom_pkg=None))
        C.top = C.main_checkout = os.path.join(self.tmp, 'repo')
        C.out = os.path.join(self.tmp, 'repo', 'out')
        text = '%s %s' % (os.path.join(C.top, 'hw', 'b.kicad_pcb'), os.path.join(C.out, 'x'))
        for _ in range(5):      # no dependence on set or hash order
            self.assertEqual(C.short(text), '<repo>/hw/b.kicad_pcb <out>/x')
        C.main_checkout = self.tmp                              # a linked worktree inside the main checkout
        self.assertEqual(C.short(os.path.join(self.tmp, 'other', 'f')), '<main-checkout>/other/f')
        self.assertEqual(C.short(os.path.join(C.top, 'f')), '<repo>/f')

    def test_check_and_quote(self):
        C = ef.Ctx(argparse.Namespace())
        self.assertTrue(C.check('a', True))
        self.assertFalse(C.check('b', False, 'why', level='warn'))
        self.assertEqual([c['level'] for c in C.checks], ['error', 'warn'])
        self.assertEqual(ef._q('a b'), '"a b"')
        self.assertEqual(ef._q(''), '""')
        self.assertEqual(ef._q('ab'), 'ab')

    def test_printed_text(self):
        self.assertEqual(ef.check_line({'name': 'zip ok', 'level': 'error', 'ok': True, 'detail': 'x'}), 'ok    zip ok')
        self.assertEqual(ef.check_line({'name': 'n', 'level': 'warn', 'ok': False, 'detail': ''}), 'WARN  n')
        self.assertEqual(ef.check_line({'name': 'n', 'level': 'error', 'ok': False, 'detail': 'why'}),
                         'FAIL  n - why')
        self.assertEqual((ef.plural(1, 'Populate 0 part'), ef.plural(2, 'page'), ef.plural(0, 'part')),
                         ('1 Populate 0 part', '2 pages', '0 parts'))
        self.assertEqual(ef.config.show_path(os.path.join('fab', 'x.txt')), 'fab/x.txt')


class UploadZipTest(TmpDirTest):
    def make_package(self, name='board_v1.0', xlsx=False):
        out = self.tmp
        base = ef.upload_base(name)
        layer_map = [{'fab_file': '%s_F_Cu.gbr' % name, 'upload_file': '%s.GTL' % base},
                     {'fab_file': '%s_Edge_Cuts.gbr' % name, 'upload_file': '%s.BOR' % base}]
        for s in layer_map:
            write(os.path.join(out, 'fab', s['fab_file']), 'G04 %s*\nM02*\n' % s['fab_file'])
        for kind in ('PTH', 'PTH_FILLED', 'NPTH'):
            write(os.path.join(out, 'fab', '%s_%s.drl' % (name, kind)), 'M48\n%\nM30\n')
        write(os.path.join(out, 'assembly', '%s.XYRS' % name),
              '# header\nC1\t1.00\t2.00\t0\t1\t1\t10.00\t20.00\t10k\tR_0603\t1\tMPN1\n'
              'C2\t3.00\t4.00\t90\t2\t1\t10.00\t20.00\t10k\tR_0603\t0\tMPN1\n')
        if xlsx:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['Designator', 'Value', 'Footprint', 'Populate', 'MPN', 'Manufacturer'])
            ws.append(['C1', '10k', 'R_0603', 1, 'MPN1', 'M'])
            ws.append(['C2', '10k', 'R_0603', 0, 'MPN1', 'M'])
            os.makedirs(os.path.join(out, 'bom'))
            wb.save(os.path.join(out, 'bom', '%s_macrofab.xlsx' % name))
        return {'name': name, 'upload_base': base, 'layer_map': layer_map, 'commit_date': '2026-01-02T03:04:06+00:00'}

    def test_upload_members_order_and_names(self):
        m = self.make_package()
        members = ef.upload_members(self.tmp, m['name'], m['layer_map'])
        self.assertEqual([arc for _, arc in members], ['board_v1_0.GTL', 'board_v1_0.BOR', 'board_v1_0_PTH.XLN',
                                                       'board_v1_0_PTH_FILLED.XLN', 'board_v1_0_NPTH.XLN',
                                                       'board_v1_0.XYRS'])

    def test_build_zip_is_flat_dated_and_byte_identical(self):
        m = self.make_package()
        C = ef.Ctx(argparse.Namespace())
        C.out = self.tmp
        out, names = ef.build_zip(C, m)
        self.assertEqual(os.path.basename(out), 'board_v1.0_macrofab_upload.zip')
        self.assertTrue(all(n.count('.') == 1 for n in names))
        with zipfile.ZipFile(out) as z:
            self.assertEqual(z.getinfo('board_v1_0.XYRS').date_time, (2026, 1, 2, 3, 4, 6))
        failed = [c for c in C.checks if not c['ok'] and c['level'] == 'error']
        self.assertEqual(failed, [])
        self.assertTrue(any('no bom' in c['name'] for c in C.checks if c['level'] == 'warn'))
        with open(out, 'rb') as f:
            first = f.read()
        os.remove(out)
        ef.build_zip(C, m)
        with open(out, 'rb') as f:
            self.assertEqual(f.read(), first)       # reproducible: fixed member dates and order

    def test_build_zip_accepts_a_utc_z_commit_date(self):
        m = self.make_package()
        m['commit_date'] = '2026-01-02T03:04:06Z'        # git's %cI for a UTC commit; Python < 3.11 rejects 'Z'
        C = ef.Ctx(argparse.Namespace())
        C.out = self.tmp
        out, _ = ef.build_zip(C, m)
        with zipfile.ZipFile(out) as z:
            self.assertEqual(z.getinfo('board_v1_0.XYRS').date_time, (2026, 1, 2, 3, 4, 6))

    def test_build_zip_with_xlsx_checks_designators(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest('openpyxl not installed')
        m = self.make_package(xlsx=True)
        C = ef.Ctx(argparse.Namespace())
        C.out = self.tmp
        out, names = ef.build_zip(C, m)
        self.assertIn('board_v1_0_macrofab.xlsx', names)
        self.assertTrue(all(c['ok'] for c in C.checks if c['name'].startswith('xlsx vs XYRS')))
        # a designator mismatch stops the zip
        xyrs = os.path.join(self.tmp, 'assembly', 'board_v1.0.XYRS')
        with open(xyrs, 'a', encoding='ascii', newline='\n') as f:
            f.write('C3\t5.00\t6.00\t0\t1\t1\t1.00\t1.00\t1u\tC_0402\t1\tMPN2\n')
        C2 = ef.Ctx(argparse.Namespace())
        C2.out = self.tmp
        with self.assertRaises(ef.StepError):
            ef.build_zip(C2, m)

    def test_zip_only_setup_errors_exit_2(self):
        mpath = os.path.join(self.tmp, 'docs', 'board_v1.0_export_manifest.json')
        for manifest, why in ((None, 'needs exactly one docs/*_export_manifest.json'),
                              ('{', 'is not a kicad-fab-export manifest'),
                              ('{"name": "board_v1.0"}', "no 'layer_map' key")):
            if manifest is not None:
                write(mpath, manifest)
            out = io.StringIO()
            with redirect_stdout(out):
                self.assertEqual(ef.main(['--zip-only', '--out', self.tmp]), 2)
            self.assertTrue(out.getvalue().startswith('ERROR: '), out.getvalue())
            self.assertIn(why, out.getvalue())

    def test_zip_only_prints_clean_check_lines(self):
        m = self.make_package()
        m.update({'checks': [], 'outputs': []})
        write(os.path.join(self.tmp, 'docs', 'board_v1.0_export_manifest.json'), json.dumps(m))
        out = io.StringIO()
        with redirect_stdout(out):
            self.assertEqual(ef.main(['--zip-only', '--out', self.tmp]), 0)
        lines = out.getvalue().splitlines()
        self.assertEqual([l for l in lines if l != l.rstrip()], [])            # no trailing spaces
        self.assertIn('WARN  upload zip: no bom/board_v1.0_macrofab.xlsx yet (run kicad-macrofab-bom, then '
                      'kicad-fab-export --zip-only)', lines)
        self.assertTrue(lines[-1].startswith('wrote ') and '\\' not in lines[-1], lines[-1])

    def test_pcbway_zip_has_the_fab_files_only(self):
        m = self.make_package()
        write(os.path.join(self.tmp, 'bom', 'board_v1.0_macrofab.xlsx'), 'the MacroFab BOM, not for PCBWay')
        C = ef.Ctx(argparse.Namespace())
        C.out = self.tmp
        out, names = ef.build_pcbway_zip(C, m, 0)
        self.assertEqual(os.path.basename(out), 'board_v1.0_pcbway_gerbers.zip')
        self.assertEqual(names, ['board_v1.0_F_Cu.gbr', 'board_v1.0_Edge_Cuts.gbr', 'board_v1.0_PTH.drl',
                                 'board_v1.0_NPTH.drl'])                   # no XYRS, no BOM, no empty filled file
        self.assertEqual([(c['name'], c['ok']) for c in C.checks],
                         [('PCBWay Gerber zip: 4 members (2 Gerbers, PTH + NPTH drills), CRC ok, byte-identical to '
                           'the fab files, no assembly or reference files', True)])
        self.assertEqual(C.notes, [])
        with zipfile.ZipFile(out) as z:
            self.assertEqual(z.getinfo('board_v1.0_PTH.drl').date_time, (2026, 1, 2, 3, 4, 6))
            self.assertEqual(z.read('board_v1.0_F_Cu.gbr'), b'G04 board_v1.0_F_Cu.gbr*\nM02*\n')
        with open(out, 'rb') as f:
            first = f.read()
        ef.build_pcbway_zip(C, m, 0)
        with open(out, 'rb') as f:
            self.assertEqual(f.read(), first)
        self.assertEqual(C.outputs, [out, out])

    def test_pcbway_zip_with_filled_vias(self):
        m = self.make_package()
        C = ef.Ctx(argparse.Namespace())
        C.out = self.tmp
        _, names = ef.build_pcbway_zip(C, m, 2)
        self.assertEqual(names[2:], ['board_v1.0_PTH.drl', 'board_v1.0_PTH_FILLED.drl', 'board_v1.0_NPTH.drl'])
        self.assertTrue(C.checks[0]['ok'])
        self.assertIn('PTH + PTH_FILLED + NPTH drills', C.checks[0]['name'])
        self.assertEqual(C.notes, ['PCBWay Gerber zip: board_v1.0_PTH_FILLED.drl holds the 2 filled-and-capped vias '
                                   '(IPC-4761 Type VII); say on the PCBWay order that the holes in that file are '
                                   'filled and capped'])
        os.remove(os.path.join(self.tmp, 'fab', 'board_v1.0_NPTH.drl'))
        with self.assertRaises(ef.StepError):
            ef.build_pcbway_zip(C, m, 2)

    def test_file_tables(self):
        self.make_package()
        C = ef.Ctx(argparse.Namespace())
        C.out = self.tmp
        mpath = write(os.path.join(self.tmp, 'docs', 'board_v1.0_export_manifest.json'), '{}')
        write(os.path.join(self.tmp, 'README.md'), 'hello')
        outs, other = ef.package_tables(C, mpath, ['assembly/board_v1.0.XYRS'])
        self.assertEqual([r['path'] for r in outs], ['assembly/board_v1.0.XYRS'])
        self.assertIn('README.md', [r['path'] for r in other])
        self.assertNotIn('docs/board_v1.0_export_manifest.json', [r['path'] for r in other])
        self.assertEqual(len(outs[0]['sha256']), 64)


def xyrs_row(ref, x, y, side=1, pop=1, kind=1, rot=0.0, value='10k', footprint='R_0603', mpn='RC0603FR-0710KL'):
    """An XYRS row as step_placement builds it (x, y in mil)."""
    return {'ref': ref, 'x': x, 'y': y, 'rot': ef.xyrs_rotation(rot, side)[1], 'side': side, 'type': kind,
            'sx': 30.0, 'sy': 60.0, 'value': value, 'footprint': footprint, 'pop': pop, 'mpn': mpn, 'kicad_rot': rot,
            'kind': 'TH' if kind == 2 else 'SMD'}


PCBWAY_ROWS = [
    xyrs_row('C1', 100.0, 200.0, value='100n', footprint='C_0402', mpn='GRM155R71C104KA88D'),
    xyrs_row('J1', 1000.0, 50.0, kind=2, rot=90.0, value='Conn_01x04', footprint='PinHeader_1x04', mpn='61300411121'),
    xyrs_row('R1', 300.0, 300.0), xyrs_row('R2', 400.0, 300.0, rot=-90.0),
    xyrs_row('R3', 500.0, 300.0, pop=0),                                   # DNP: never on the fitted line
    xyrs_row('R4', 600.0, 300.0),                                          # another manufacturer spelling
    xyrs_row('R10', 700.0, 300.0, side=2, rot=90.0),
    xyrs_row('TP1', 800.0, 0.0, pop=0, value='TestPoint', footprint='TestPoint_Pad_D1.0mm', mpn=''),
    xyrs_row('U1', 0.0, 1000.0, side=2, rot=45.5, value='STM32', footprint='QFN-48', mpn='STM32L431CCU6'),
]
PCBWAY_MFR = {'C1': 'Murata', 'J1': 'Wurth Elektronik', 'R1': 'Yageo', 'R2': 'Yageo', 'R3': 'Yageo',
              'R4': 'YAGEO', 'R10': 'Yageo', 'TP1': '', 'U1': 'STMicroelectronics'}


def write_xyrs(path, rows):
    """An XYRS file with these rows, in the layout step_placement writes (one header line)."""
    lines = ['#Designator\tX-Loc\tY-Loc\tRotation\tSide\tType\tX-Size\tY-Size\tValue\tFootprint\tPopulate\tMPN']
    lines += ['\t'.join([r['ref'], '%.2f' % r['x'], '%.2f' % r['y'], '%d' % r['rot'], str(r['side']), str(r['type']),
                         '%.2f' % r['sx'], '%.2f' % r['sy'], r['value'], r['footprint'], str(r['pop']), r['mpn']])
              for r in rows]
    return write(path, '\n'.join(lines) + '\n')


class PcbwayRulesTest(unittest.TestCase):
    def test_type_comes_from_the_xyrs_row(self):
        self.assertEqual([ef.pcbway_type(r) for r in PCBWAY_ROWS],
                         ['SMD', 'thru-hole', 'SMD', 'SMD', 'DNS', 'SMD', 'SMD', 'DNS', 'SMD'])

    def test_bom_lines_group_and_order(self):
        lines = ef.pcbway_bom_lines(PCBWAY_ROWS, PCBWAY_MFR)
        self.assertTrue(all(list(l) == ef.PCBWAY_BOM_COLUMNS for l in lines))
        got = [(l['Item #'], l['*Designator'], l['*Qty'], l['Manufacturer'], l['Type']) for l in lines]
        self.assertEqual(got, [(1, 'C1', 1, 'Murata', 'SMD'),
                               (2, 'J1', 1, 'Wurth Elektronik', 'thru-hole'),
                               (3, 'R1, R2, R10', 3, 'Yageo', 'SMD'),          # natural order, both sides
                               (4, 'R4', 1, 'YAGEO', 'SMD'),                   # the manufacturer is in the key
                               (5, 'U1', 1, 'STMicroelectronics', 'SMD'),
                               (6, 'R3', 1, 'Yageo', 'DNS'),                   # DNS lines last
                               (7, 'TP1', 1, '', 'DNS')])
        r = lines[2]
        self.assertEqual((r['*Mfg Part #'], r['Description / Value'], r['*Package/Footprint'],
                          r['Your Instructions / Notes']), ('RC0603FR-0710KL', '10k', 'R_0603', ''))

    def test_centroid_rows(self):
        self.assertEqual(ef.pcbway_centroid(PCBWAY_ROWS, 1),
                         [['C1', '2.5400', '5.0800', 'Top', '0'], ['J1', '25.4000', '1.2700', 'Top', '90'],
                          ['R1', '7.6200', '7.6200', 'Top', '0'], ['R2', '10.1600', '7.6200', 'Top', '270'],
                          ['R4', '15.2400', '7.6200', 'Top', '0']])
        # bottom: KiCad's own rotation (the XYRS says 270 for R10 and 225 for U1), X not mirrored
        self.assertEqual(ef.pcbway_centroid(PCBWAY_ROWS, 2),
                         [['R10', '17.7800', '7.6200', 'Bottom', '90'], ['U1', '0.0000', '25.4000', 'Bottom', '45.5']])

    def test_number_formats(self):
        for ori, want in ((90, '90'), (-90, '270'), (45.5, '45.5'), (359.9996, '0'), (180.0001, '180'), (-0.0, '0'),
                          (0.1234, '0.123'), (720, '0')):
            self.assertEqual(ef.pcbway_rotation(ori), want, ori)
        self.assertEqual((ef.pcbway_mm(1000), ef.pcbway_mm(-1e-9), ef.pcbway_mm(0.004)), ('25.4000', '0.0000',
                                                                                         '0.0001'))

    def test_step_needs_openpyxl(self):
        with mock.patch.object(ef.make_macrofab_xlsx, 'openpyxl', None):
            with self.assertRaises(ef.StepError) as e:
                ef.step_pcbway(ef.Ctx(argparse.Namespace()))
        self.assertIn('needs openpyxl', str(e.exception))

    def test_zip_members(self):
        layer_map = [{'fab_file': 'b_v1_F_Cu.gbr'}, {'fab_file': 'b_v1_Edge_Cuts.gbr'}]
        names = [n for _, n in ef.pcbway_zip_members('out', 'b_v1', layer_map, 0)]
        self.assertEqual(names, ['b_v1_F_Cu.gbr', 'b_v1_Edge_Cuts.gbr', 'b_v1_PTH.drl', 'b_v1_NPTH.drl'])
        members = ef.pcbway_zip_members('out', 'b_v1', layer_map, 3)
        self.assertEqual([n for _, n in members][2:], ['b_v1_PTH.drl', 'b_v1_PTH_FILLED.drl', 'b_v1_NPTH.drl'])
        self.assertEqual(members[0][0], os.path.join('out', 'fab', 'b_v1_F_Cu.gbr'))


@unittest.skipIf(ef.make_macrofab_xlsx.openpyxl is None, 'openpyxl not installed')
class StepPcbwayTest(TmpDirTest):
    """step_pcbway on a synthetic package: the files, and the read-back checks against the XYRS and the pos file."""

    def ctx(self, rows=PCBWAY_ROWS, pos_rot=None):
        C = ef.Ctx(argparse.Namespace())
        C.out, C.prefix, C.sha, C.generator = self.tmp, 'board_v1.0', 'abc123', 'export_fab.py'
        C.commit_date = '2026-01-02T03:04:06+00:00'
        C.xyrs_rows = rows
        C.xyrs_path = write_xyrs(os.path.join(self.tmp, 'assembly', 'board_v1.0.XYRS'), rows)
        C.comps = {r['ref']: {'fields': {'Manufacturer': PCBWAY_MFR[r['ref']]}} for r in rows}
        C.comps['J1'] = {'fields': {'MANUFACTURER': ' Wurth Elektronik '}}       # the other spelling, stripped
        C.counts = {'fitted_top': sum(1 for r in rows if r['pop'] and r['side'] == 1),
                    'fitted_bottom': sum(1 for r in rows if r['pop'] and r['side'] == 2)}
        # KiCad's pos file: raw rotation, DNP parts left out (here: R3; TP1 is in it, excluded from the BOM only)
        pos_rot = pos_rot or {}
        C.pos_rows = [{'Ref': r['ref'], 'Rot': '%f' % pos_rot.get(r['ref'], r['kicad_rot']),
                       'Side': 'top' if r['side'] == 1 else 'bottom'} for r in rows if r['ref'] != 'R3']
        return C

    def failed(self, C):
        return [(c['name'], c['detail']) for c in C.checks if not c['ok']]

    def test_files_and_checks(self):
        C = self.ctx()
        detail = ef.step_pcbway(C)
        self.assertEqual(detail, 'BOM 7 lines (2 DNS) for 9 XYRS rows; centroid 5 top / 2 bottom')
        self.assertEqual(self.failed(C), [])
        self.assertEqual([c['name'].split(':')[0] for c in C.checks], ['PCBWay BOM', 'PCBWay centroid',
                                                                       'PCBWay centroid vs KiCad pos'])
        asm = os.path.join(self.tmp, 'assembly')
        with open(os.path.join(asm, 'board_v1.0_pcbway_bottom.csv'), encoding='ascii', newline='') as f:
            self.assertEqual(f.read(), 'Designator,Mid X(mm),Mid Y(mm),Layer,Rotation\n'
                                       'R10,17.7800,7.6200,Bottom,90\nU1,0.0000,25.4000,Bottom,45.5\n')
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(asm, 'board_v1.0_pcbway_bom.xlsx'))
        ws = wb.worksheets[0]
        got = [[c.value for c in r] for r in ws.iter_rows()]
        self.assertEqual(got[0], ef.PCBWAY_BOM_COLUMNS)
        self.assertEqual(got[2], [2, 'J1', 1, 'Wurth Elektronik', '61300411121', 'Conn_01x04', 'PinHeader_1x04',
                                  'thru-hole', None])
        self.assertEqual(got[-1], [7, 'TP1', 1, None, None, 'TestPoint', 'TestPoint_Pad_D1.0mm', 'DNS', None])
        self.assertEqual((wb.properties.creator, wb.properties.title), ('export_fab.py', 'board_v1.0 PCBWay BOM'))
        self.assertEqual(C.pcbway['centroid']['bottom'], {'path': 'assembly/board_v1.0_pcbway_bottom.csv', 'rows': 2})
        self.assertEqual((C.pcbway['bom_lines'], C.pcbway['dns_lines']), (7, 2))
        self.assertEqual(sorted(os.path.relpath(p, self.tmp).replace(os.sep, '/') for p in C.outputs),
                         ['assembly/board_v1.0_pcbway_bom.xlsx', 'assembly/board_v1.0_pcbway_bottom.csv',
                          'assembly/board_v1.0_pcbway_top.csv'])

    def test_reproducible(self):
        path = os.path.join(self.tmp, 'assembly', 'board_v1.0_pcbway_bom.xlsx')
        ef.step_pcbway(self.ctx())
        with open(path, 'rb') as f:
            first = f.read()
        ef.step_pcbway(self.ctx())
        with open(path, 'rb') as f:
            self.assertEqual(f.read(), first)
        with zipfile.ZipFile(path) as z:
            self.assertTrue(all(i.date_time == (2026, 1, 2, 3, 4, 6) for i in z.infolist()))

    def test_a_bottom_rotation_with_180_added_fails(self):
        C = self.ctx(pos_rot={'R10': 270.0})         # as if the centroid used the XYRS bottom rule
        ef.step_pcbway(C)
        failed = self.failed(C)
        self.assertEqual(len(failed), 1)
        self.assertTrue(failed[0][0].startswith('PCBWay centroid vs KiCad pos: rotation = pos Rot'), failed)
        self.assertEqual(failed[0][1], 'R10 pos 270.000000 bottom -> centroid 90 bottom')

    def test_files_that_disagree_with_the_xyrs_fail(self):
        C = self.ctx()
        rows = [dict(r) for r in PCBWAY_ROWS]
        rows[2]['value'] = '1k'                      # R1 in the XYRS file differs from what the BOM and csv carry
        rows[3]['x'] = 401.0
        write_xyrs(C.xyrs_path, rows + [xyrs_row('R99', 1.0, 1.0)])
        ef.step_pcbway(C)
        failed = dict(self.failed(C))
        bom = [d for n, d in failed.items() if n.startswith('PCBWay BOM')][0]
        self.assertIn("only XYRS ['R99']", bom)
        self.assertIn('R1: line has Type/MPN/Value/Footprint SMD/RC0603FR-0710KL/10k/R_0603, XYRS '
                      'SMD/RC0603FR-0710KL/1k/R_0603', bom)
        centroid = [n for n in failed if n.startswith('PCBWay centroid: ')][0]
        self.assertIn('max diff 0.0254 mm', centroid)
        self.assertIn("top designators ['C1', 'J1', 'R1', 'R2', 'R4']", failed[centroid])


class _FakePage:
    def __init__(self, words):
        self.words = words

    def get_text(self, kind):
        assert kind == 'words'
        return [(0, 0, 1, 1, w, 0, 0, i) for i, w in enumerate(self.words)]


class _FakeDoc:
    def __init__(self, path):
        with open(path, encoding='ascii') as f:
            self.pages = [_FakePage(line.split()) for line in f.read().splitlines()[1:]]
        self.page_count = len(self.pages)

    def __getitem__(self, i):
        return self.pages[i]

    def close(self):
        pass


class AssemblyPdfTest(TmpDirTest):
    """step_assembly_pdf without pypdf: two one-page files, and the designator check still runs on them."""

    def ctx(self, printed):
        C = ef.Ctx(argparse.Namespace())
        C.out, C.work, C.prefix, C.board = self.tmp, self.tmp, 'board_v1.0', 'board.kicad_pcb'
        os.makedirs(os.path.join(self.tmp, 'assembly'))     # main() creates fab/, assembly/ and docs/
        C.xyrs_rows = [{'ref': 'R1', 'side': 1}, {'ref': 'C1', 'side': 1}, {'ref': 'JP1', 'side': 2}]

        def kcli(*args):        # a stand-in "PDF" per side: a header line, then the words printed on the page
            layer = args[args.index('-l') + 1].split(',')[0]
            write(args[args.index('-o') + 1], '%%PDF-1.4\n%s\n' % ' '.join(printed[layer]))
            return ''
        C.kcli = kcli
        return C

    def run_step(self, printed, pymupdf=True):
        fake = types.ModuleType('pymupdf')
        fake.open = _FakeDoc
        mods = {'pypdf': None, 'pymupdf': fake if pymupdf else None, 'fitz': None}
        C = self.ctx(printed)
        with mock.patch.dict(sys.modules, mods):
            detail = ef.step_assembly_pdf(C)
        return C, detail

    def test_no_pypdf_still_checks_designators(self):
        C, detail = self.run_step({'F.Fab': ['R1', 'C1', 'Title'], 'B.Fab': ['JP1']})
        for side in ('top', 'bottom'):
            self.assertTrue(os.path.isfile(os.path.join(self.tmp, 'assembly', 'board_v1.0_assembly_%s.pdf' % side)))
        by_name = {c['name']: c for c in C.checks}
        self.assertFalse(by_name['assembly PDF: pypdf missing, written as two files']['ok'])
        self.assertTrue(by_name["assembly PDF: every XYRS designator is printed on its side's page (3)"]['ok'])
        self.assertEqual(detail, 'two files (no pypdf), all 3 designators found')

    def test_no_pypdf_missing_designator_fails(self):
        C, detail = self.run_step({'F.Fab': ['R1', 'JP1'], 'B.Fab': ['C1']})
        c = [c for c in C.checks if c['name'].startswith('assembly PDF: every XYRS designator')][0]
        self.assertEqual((c['ok'], c['level'], c['detail']), (False, 'error', 'missing C1 JP1'))
        self.assertTrue(detail.endswith(', 2 missing'))

    def test_no_pypdf_no_pymupdf_warns_twice(self):
        C, detail = self.run_step({'F.Fab': ['R1', 'C1'], 'B.Fab': ['JP1']}, pymupdf=False)
        self.assertEqual([(c['name'], c['level'], c['ok']) for c in C.checks],
                         [('assembly PDF: pypdf missing, written as two files', 'warn', False),
                          ('assembly PDF: designators not checked (no PyMuPDF)', 'warn', False)])
        self.assertEqual(detail, 'two files (no pypdf)')


class ExpectKeysTest(unittest.TestCase):
    def test_keys(self):
        self.assertEqual(collections.Counter(ef.EXPECT_KEYS)['fitted'], 1)
        self.assertIn('filled', ef.EXPECT_KEYS)


if __name__ == '__main__':
    unittest.main()

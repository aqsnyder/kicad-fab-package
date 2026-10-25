import datetime
import importlib.util
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

from kicad_fab_package import make_macrofab_xlsx as mx

from ._util import rmtree

try:
    import openpyxl
except ImportError:
    openpyxl = None

HAVE_LXML = importlib.util.find_spec('lxml') is not None
# the folder that holds the kicad_fab_package these tests imported, for a fresh interpreter to import the same code
PKG_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(mx.__file__)))


def fresh_python(code, *args, **env_set):
    """Run code in a new interpreter (openpyxl picks its XML writer once, on its first import), with OPENPYXL_LXML
    unset unless given. Returns stdout."""
    env = {k: v for k, v in os.environ.items() if k != 'OPENPYXL_LXML'}
    env.update(env_set)
    env['PYTHONPATH'] = os.pathsep.join([PKG_ROOT] + ([env['PYTHONPATH']] if env.get('PYTHONPATH') else []))
    env['PYTHONDONTWRITEBYTECODE'] = '1'
    r = subprocess.run([sys.executable, '-c', code] + list(args), capture_output=True, text=True, env=env,
                       cwd=PKG_ROOT)
    if r.returncode != 0:
        raise AssertionError('fresh interpreter failed:\n%s%s' % (r.stdout, r.stderr))
    return r.stdout.strip()


def comp(ref, value, footprint, mpn='', manufacturer='', dnp=False, excl_bom=False, excl_board=False, **fields):
    f = {'MPN': mpn, 'Manufacturer': manufacturer}
    f.update(fields)
    return {'ref': ref, 'value': value, 'footprint': footprint, 'fields': f, 'dnp': dnp, 'excl_bom': excl_bom,
            'excl_board': excl_board}


def bomrow(ref, value, footprint, mpn, manufacturer='', MANUFACTURER='', dnp='', excl=''):
    return {'Reference': ref, 'Value': value, 'Footprint': footprint, 'MPN': mpn, 'Manufacturer': manufacturer,
            'MANUFACTURER': MANUFACTURER, 'DNP': dnp, 'ExclBOM': excl, 'ExclBoard': ''}


def fp(ref, lib, value, mpn='', layer='F.Cu', attrs=()):
    return {'ref': ref, 'lib': lib, 'layer': layer, 'value': value, 'mpn': mpn, 'attrs': set(attrs)}


def design():
    """A small synthetic design: two fitted resistors, a DNP strap of the same value and package, a test point
    excluded from the BOM, a fiducial, a board-only logo, and a part on the bottom side."""
    comps = {
        'R1': comp('R1', '0R', 'Resistor_SMD:R_0603_1608Metric', 'R0603-0R-TEST', 'Murata'),
        'R2': comp('R2', '0R', 'Resistor_SMD:R_0603_1608Metric', 'R0603-0R-TEST', 'Murata Electronics'),
        'R3': comp('R3', '0R', 'Resistor_SMD:R_0603_1608Metric', 'R0603-0R-TEST', 'Murata', dnp=True,
                   excl_bom=True),
        'C10': comp('C10', '100n', 'Capacitor_SMD:C_0402_1005Metric', 'MF-TEST-0001', 'MacroFab'),
        'C2': comp('C2', '100n', 'Capacitor_SMD:C_0402_1005Metric', 'C0402-100N-TEST', '', MANUFACTURER='Murata'),
        'TP1': comp('TP1', 'TP', 'TestPoint:TestPoint_Pad_D1.0mm', excl_bom=True),
        'FID1': comp('FID1', 'Fiducial', 'Fiducial:Fiducial_1mm_Mask2mm', excl_bom=True),
        'U1': comp('U1', 'LUX1', 'Package_SON:USON-6', 'ABC1001DNPR', 'Acme Semi'),
    }
    bom = [bomrow('R1', '0R', 'Resistor_SMD:R_0603_1608Metric', 'R0603-0R-TEST', 'Murata'),
           bomrow('R2', '0R', 'Resistor_SMD:R_0603_1608Metric', 'R0603-0R-TEST', 'Murata Electronics'),
           bomrow('C10', '100n', 'Capacitor_SMD:C_0402_1005Metric', 'MF-TEST-0001', 'MacroFab'),
           bomrow('C2', '100n', 'Capacitor_SMD:C_0402_1005Metric', 'C0402-100N-TEST', '', 'Murata'),
           bomrow('U1', 'LUX1', 'Package_SON:USON-6', 'ABC1001DNPR', 'Acme Semi')]
    fps = [fp('R1', 'Resistor_SMD:R_0603_1608Metric', '0R', 'R0603-0R-TEST'),
           fp('R2', 'Resistor_SMD:R_0603_1608Metric', '0R', 'R0603-0R-TEST', layer='B.Cu'),
           fp('R3', 'Resistor_SMD:R_0603_1608Metric', '0R', 'R0603-0R-TEST', attrs=('dnp', 'exclude_from_bom')),
           fp('C10', 'Capacitor_SMD:C_0402_1005Metric', '100n', 'MF-TEST-0001'),
           fp('C2', 'Capacitor_SMD:C_0402_1005Metric', '100n', 'C0402-100N-TEST'),
           fp('TP1', 'TestPoint:TestPoint_Pad_D1.0mm', 'TP', attrs=('exclude_from_bom',)),
           fp('FID1', 'Fiducial:Fiducial_1mm_Mask2mm', 'Fiducial', attrs=('exclude_from_bom',)),
           fp('U1', 'Package_SON:USON-6', 'LUX1', 'ABC1001DNPR'),
           fp('LOGO1', 'Logo:Logo', 'LOGO', attrs=('board_only',))]
    return bom, comps, fps


class RefsTest(unittest.TestCase):
    def test_natkey_order(self):
        self.assertEqual(sorted(['C10', 'C2', 'R1', 'C1', 'TP3'], key=mx.natkey), ['C1', 'C2', 'C10', 'R1', 'TP3'])

    def test_expand_and_compress(self):
        self.assertEqual(mx.expand_refs('TP1-TP3, J5,R5-7'), {'TP1', 'TP2', 'TP3', 'J5', 'R5', 'R6', 'R7'})
        self.assertEqual(mx.expand_refs(''), set())
        self.assertEqual(mx.compress_refs(['C1', 'C2', 'C3', 'C7', 'R1']), 'C1-C3, C7, R1')
        self.assertEqual(mx.compress_refs([]), '')

    def test_small_helpers(self):
        self.assertEqual(mx.fp_name('Lib:Name'), 'Name')
        self.assertEqual(mx.fp_lib('Lib:Name'), 'Lib')
        self.assertEqual(mx.fp_lib('Name'), '')
        self.assertEqual(mx.unq('"a\\"b"'), 'a"b')
        self.assertEqual(mx.md('a|b\nc'), 'a\\|b c')
        self.assertEqual(mx.shell_quote('a b'), "'a b'")
        self.assertEqual(mx.shell_quote('ab'), 'ab')
        self.assertEqual(mx.norm_mpn(' rc-0603 '), 'RC0603')


class BuildRowsTest(unittest.TestCase):
    def test_rows_and_populate(self):
        bom, comps, fps = design()
        rows, res = mx.build_rows(bom, comps, [], fps)
        self.assertEqual(res['errors'], [])
        self.assertEqual([r['Designator'] for r in rows], ['C2', 'C10', 'R1', 'R2', 'R3', 'TP1', 'U1'])
        pop = {r['Designator']: r['Populate'] for r in rows}
        self.assertEqual(pop, {'C2': 1, 'C10': 1, 'R1': 1, 'R2': 1, 'R3': 0, 'TP1': 0, 'U1': 1})
        by = {r['Designator']: r for r in rows}
        self.assertEqual(by['R3']['MPN'], 'R0603-0R-TEST')          # a DNP row keeps its MPN
        self.assertEqual(by['C2']['Manufacturer'], 'Murata')          # MANUFACTURER fallback
        self.assertEqual(res['fallback'], ['C2'])
        self.assertEqual(by['R1']['Footprint'], 'R_0603_1608Metric')
        self.assertEqual([o[0] for o in res['omitted']], ['FID1'])
        self.assertEqual(res['omitted'][0][2], 'Fiducial footprint: not a part')   # no baseline configured
        self.assertEqual([f['ref'] for f in res['board_only']], ['LOGO1'])
        self.assertEqual(res['side'][(1, 'B.Cu')], 1)

    def test_baseline_label_words_the_omitted_reason(self):
        bom, comps, fps = design()
        _, res = mx.build_rows(bom, comps, [], fps, baseline_label='v1.1')
        self.assertEqual(res['omitted'][0][2], 'Fiducial footprint: not a part (no row in the as-built v1.1 BOM)')

    def test_non_part_rule_settings(self):
        import re
        bom, comps, fps = design()
        # a fiducial from a custom library: found by designator by default, kept as a row with no designator rule
        comps['FID1']['footprint'] = 'MyMech:Fid_1mm'
        fps[6]['lib'] = 'MyMech:Fid_1mm'
        _, res = mx.build_rows(bom, comps, [], fps)
        self.assertEqual([(o[0], o[2]) for o in res['omitted']], [('FID1', 'fiducial / mounting-hole designator: '
                                                                            'not a part')])
        rows, res = mx.build_rows(bom, comps, [], fps, non_part_refs=None)
        self.assertEqual(res['omitted'], [])
        self.assertIn('FID1', [r['Designator'] for r in rows])
        _, res = mx.build_rows(bom, comps, [], fps, non_part_refs=None, non_part_libs=('MyMech',))
        self.assertEqual([o[0] for o in res['omitted']], ['FID1'])
        _, res = mx.build_rows(bom, comps, [], fps, non_part_refs=re.compile(r'TP\d+$'), non_part_libs=())
        self.assertEqual([o[0] for o in res['omitted']], ['TP1'])      # Populate 0 and no MPN

    def test_errors_are_collected(self):
        bom, comps, fps = design()
        bom.append(bomrow('R3', '0R', 'Resistor_SMD:R_0603_1608Metric', 'R0603-0R-TEST', dnp='DNP'))
        fps[0]['mpn'] = 'OTHER'
        comps['R9?'] = comp('R9?', '1k', 'Resistor_SMD:R_0603_1608Metric', 'X')
        fps.append(fp('R9?', 'Resistor_SMD:R_0603_1608Metric', '1k', 'X'))
        comps['C3'] = comp('C3', '1u', 'Capacitor_SMD:C_0402_1005Metric', '')   # fitted, no MPN, not on the board
        rows, res = mx.build_rows(bom, comps, ['R1'], fps)
        text = '\n'.join(res['errors'])
        self.assertIn('duplicate designator R1 in the netlist', text)
        self.assertIn('R3: DNP but still in the BOM', text)
        self.assertIn('R1: MPN differs: board', text)
        self.assertIn('unannotated symbol R9?', text)
        self.assertIn('C3: symbol has no footprint on the board', text)
        self.assertIn('C3: in the BOM per its flags but missing from the BOM export', text)
        self.assertIn('C3: Populate 1 without an MPN', text)

    def test_whitespace_is_stripped_and_reported(self):
        bom, comps, fps = design()
        bom[0]['Value'] = '0R '
        comps['R1']['value'] = '0R '
        rows, res = mx.build_rows(bom, comps, [], fps)
        self.assertEqual(res['stripped'], [('R1', 'Value', '0R ')])
        self.assertEqual([r for r in rows if r['Designator'] == 'R1'][0]['Value'], '0R')


class HygieneTest(unittest.TestCase):
    def test_flags(self):
        bom, comps, fps = design()
        rows, _ = mx.build_rows(bom, comps, [], fps)
        h = mx.hygiene(rows)
        self.assertEqual(h['mixed_populate'], [(('0R', 'R_0603_1608Metric'), ['R1', 'R2'], ['R3'])])
        self.assertEqual(h['dnp_text'], [('U1', 'MPN', 'ABC1001DNPR')])
        self.assertEqual(h['house'], [(('MF-TEST-0001', 'MacroFab'), ['C10'])])
        self.assertEqual(h['mfr_variants'], [['Murata', 'Murata Electronics']])
        self.assertEqual(h['mpn_two_mfrs'], [('R0603-0R-TEST', ['Murata', 'Murata Electronics'])])
        self.assertEqual(h['collisions'], [(('100n', 'C_0402_1005Metric'),
                                            [('C0402-100N-TEST', ['C2']), ('MF-TEST-0001', ['C10'])])])
        self.assertEqual(h['pop0_no_mpn'], ['TP1'])
        self.assertEqual(h['odd_mpn'], [])

    def test_odd_mpn_and_formula(self):
        rows = [{'Designator': 'F1', 'Value': '=1', 'Footprint': 'F', 'Populate': 1, 'MPN': 'PTC100/24X',
                 'Manufacturer': 'CUI Devices'}]
        h = mx.hygiene(rows)
        self.assertEqual(h['odd_mpn'], [('F1', 'PTC100/24X', '/')])
        self.assertEqual(h['formula_like'], [('F1', 'Value', '=1')])
        self.assertEqual(h['mfr_known'][0][0], 'CUI Devices')


class ReportsTest(unittest.TestCase):
    def test_diff_and_checks_reports_use_the_labels(self):
        bom, comps, fps = design()
        rows, res = mx.build_rows(bom, comps, [], fps, baseline_label='v1.1')
        base = {'R1': {'Designator': 'R1', 'Value': '0R', 'Footprint': 'R_0603_1608Metric', 'Populate': 1,
                       'MPN': 'R0603-0R-TEST', 'Manufacturer': 'Murata'},
                'R4': {'Designator': 'R4', 'Value': '1k', 'Footprint': 'R_0603_1608Metric', 'Populate': 1,
                       'MPN': 'OLD', 'Manufacturer': 'X'},
                'U1': {'Designator': 'U1', 'Value': 'LUX1', 'Footprint': 'USON-6', 'Populate': 0,
                       'MPN': 'ABC1001DNPR', 'Manufacturer': 'Acme Semi'}}
        text, info = mx.diff_md('board_v1.2', 'commit `abc`', rows, 'abc:board/old.xlsx', base, [], [],
                                label='v1.1', new_label='v1.2', generator='tools/make_macrofab_xlsx.py')
        self.assertTrue(text.startswith('# board_v1.2 BOM vs the as-built v1.1 BOM\n'))
        self.assertIn('Generated by `kicad-macrofab-bom` (`tools/make_macrofab_xlsx.py`); do not edit by hand.', text)
        self.assertIn('(v1.1 -> v1.2)', text)
        self.assertIn('## v1.1 fitted MPNs no longer on the BOM (1)', text)
        self.assertEqual(info['removed'], ['R4'])
        self.assertIn('U1', [c for c in text.split('## Changed')[1].split('\n') if c.startswith('| U1')][0])
        rows[0]['MPN'] = 'GRM155/ODD'                 # an odd MPN brings the as-built comparison column
        h = mx.hygiene(rows)
        md = mx.checks_md('board_v1.2', 'commit `abc`', '10.0.7', ['kicad-cli ...'], rows, res, h, [],
                          base=base, refs=[], label='v1.1')
        self.assertIn('Generated by `kicad-macrofab-bom` (`make_macrofab_xlsx.py`); do not edit by hand.', md)
        self.assertNotIn('kproject.py', md)
        self.assertIn('As-built v1.1', md)
        self.assertIn('this commit is on no branch or tag', md)
        self.assertIn('(not parts; the as-built v1.1 BOM has none)', md)
        md = mx.checks_md('board_v1.2', 'commit `abc`', '10.0.7', ['kicad-cli ...'], rows, res, h, [], refs=['x'])
        self.assertIn('- No row: fiducials and mounting holes (not parts), board-only', md)    # no baseline
        self.assertIn('| As-built baseline |', md)
        self.assertNotIn('no branch or tag', md)

    def test_grouping_advice_states_no_macrofab_rule(self):
        # how MacroFab groups BOM lines is MacroFab's to change: the report flags the facts and points at the
        # walkthrough, without a grouping rule or a button the walkthrough does not describe
        bom, comps, fps = design()
        rows, res = mx.build_rows(bom, comps, [], fps)
        h = mx.hygiene(rows)
        self.assertTrue(h['collisions'] and h['mixed_populate'])
        md = mx.checks_md('board_v1.2', 'commit `abc`', '10.0.7', ['kicad-cli ...'], rows, res, h, [], refs=['x'])
        self.assertIn('### Value + footprint collisions (same Value and footprint, different MPNs) (1)\n', md)
        self.assertIn('### Value + footprint groups that mix fitted and Populate 0 rows (1)\n', md)
        for claim in ('groups BOM lines by', 'grouped by', 'Ungroup', 'New Group', 'merged the line'):
            self.assertFalse(claim in md, 'the report still says %r' % claim)
        pointer = '(see "Line grouping and Populate", section 7 of the MacroFab walkthrough, docs/macrofab.md in ' \
                  'kicad-fab-package).'
        self.assertIn('check that each of these designators keeps its own MPN ' + pointer, md)
        self.assertIn('the fitted ones must stay fitted ' + pointer, md)

    def test_no_baseline_configured(self):
        self.assertEqual(mx.load_baseline(None, '.', 'abc', 'board'), (None, None, [], []))


@unittest.skipIf(openpyxl is None, 'openpyxl not installed')
class XlsxTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='kfp_xlsx_')

    def tearDown(self):
        rmtree(self.tmp)

    def test_shared_strings_absent_cells_and_reproducible(self):
        rows = [{'Designator': 'C1', 'Value': '100n', 'Footprint': 'C_0402', 'Populate': 1, 'MPN': 'GRM',
                 'Manufacturer': 'Murata'},
                {'Designator': 'R1', 'Value': '=1k', 'Footprint': 'R_0603', 'Populate': 0, 'MPN': '',
                 'Manufacturer': ' lead space'}]
        when = datetime.datetime(2026, 1, 2, 3, 4, 6)
        a, b = os.path.join(self.tmp, 'a.xlsx'), os.path.join(self.tmp, 'b.xlsx')
        mx.write_xlsx(a, rows, when, 'T', 'D')
        mx.write_xlsx(b, rows, when, 'T', 'D')
        with open(a, 'rb') as f1, open(b, 'rb') as f2:
            self.assertEqual(f1.read(), f2.read())
        with zipfile.ZipFile(a) as z:
            self.assertIn('xl/sharedStrings.xml', z.namelist())
            sheet = z.read('xl/worksheets/sheet1.xml')
            self.assertNotIn(b'inlineStr', sheet)
            self.assertIn(b'<t xml:space="preserve"> lead space</t>', z.read('xl/sharedStrings.xml'))
            self.assertIn(b'2026-01-02T03:04:06Z', z.read('docProps/core.xml'))
            self.assertTrue(all(i.date_time == (2026, 1, 2, 3, 4, 6) for i in z.infolist()))
        ws = openpyxl.load_workbook(a).worksheets[0]
        got = [[c.value for c in r] for r in ws.iter_rows()]
        self.assertEqual(got[0], mx.COLUMNS)
        self.assertEqual(got[1], ['C1', '100n', 'C_0402', 1, 'GRM', 'Murata'])
        self.assertEqual(got[2], ['R1', '=1k', 'R_0603', 0, None, ' lead space'])   # text, never a formula
        self.assertEqual(ws['B3'].data_type, 's')

    def test_load_baseline_from_file(self):
        p = os.path.join(self.tmp, 'base.xlsx')
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Designator', 'Value', 'Footprint', 'Populate', 'MPN', 'Manufacturer', 'MPN', 'Manufacturer'])
        ws.append(['C1', '100n', 'C_0402', 1, 'GRM', 'Murata', 'ALT1', 'AltCo'])
        ws.append(['C1', '100n', 'C_0402', 'yes', 'GRM', 'Murata', None, None])
        ws.append(['R1', '1k', 'R_0603', 0.0, 'RC', 'Yageo', None, None])
        wb.save(p)
        label, base, alts, dups = mx.load_baseline(p, '.', 'abc', 'board')
        self.assertEqual(label, p)
        self.assertEqual(base['C1']['Populate'], 1)
        self.assertEqual(base['R1']['Populate'], 0)
        self.assertEqual(alts, [('C1', 'ALT1', 'AltCo')])
        self.assertEqual(dups, ['C1'])


# writes the rows of XlsxTest to argv[1]; argv[2] == 'no-lxml' hides lxml as if it were not installed
WRITE_XLSX = '''
import datetime, sys
if sys.argv[2] == 'no-lxml':
    sys.modules['lxml'] = None
from kicad_fab_package import make_macrofab_xlsx as mx
import openpyxl.xml
rows = [{'Designator': 'C1', 'Value': '100n', 'Footprint': 'C_0402', 'Populate': 1, 'MPN': 'GRM',
         'Manufacturer': 'Murata'},
        {'Designator': 'R1', 'Value': '=1k', 'Footprint': 'R_0603', 'Populate': 0, 'MPN': '',
         'Manufacturer': ' lead space'}]
mx.write_xlsx(sys.argv[1], rows, datetime.datetime(2026, 1, 2, 3, 4, 6), 'T', 'D')
print(openpyxl.xml.LXML)
'''


@unittest.skipIf(openpyxl is None, 'openpyxl not installed')
class XmlWriterTest(unittest.TestCase):
    """openpyxl's lxml and standard-library writers give different bytes for the same cells, so the tool picks the
    standard-library one unless OPENPYXL_LXML says otherwise."""

    def test_standard_library_writer_unless_the_environment_chooses(self):
        code = ('import os\nimport kicad_fab_package.make_macrofab_xlsx\nimport openpyxl.xml as x\n'
                'print(os.environ.get("OPENPYXL_LXML"), x.LXML, x.lxml_available())')
        self.assertEqual(fresh_python(code).split()[:2], ['False', 'False'])
        env, lxml, available = fresh_python(code, OPENPYXL_LXML='True').split()
        self.assertEqual(env, 'True')                   # an explicit choice is kept
        self.assertEqual(lxml, available)               # lxml when it is installed

    @unittest.skipUnless(HAVE_LXML, 'lxml not installed: nothing to compare')
    def test_same_bytes_with_and_without_lxml(self):
        tmp = tempfile.mkdtemp(prefix='kfp_lxml_')
        try:
            out, writer = {}, {}
            for mode in ('lxml', 'no-lxml'):
                p = os.path.join(tmp, mode + '.xlsx')
                writer[mode] = fresh_python(WRITE_XLSX, p, mode)                # openpyxl.xml.LXML
                with open(p, 'rb') as f:
                    out[mode] = f.read()
            self.assertEqual(len(out['lxml']), len(out['no-lxml']))
            self.assertEqual(out['lxml'], out['no-lxml'])
            self.assertEqual(writer, {'lxml': 'False', 'no-lxml': 'False'})
        finally:
            rmtree(tmp)


class SourcesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='kfp_src_')

    def tearDown(self):
        rmtree(self.tmp)

    def test_board_footprints(self):
        board = ('(kicad_pcb\n\t(version 20250000)\n'
                 '\t(footprint "Resistor_SMD:R_0603"\n\t\t(layer "B.Cu")\n'
                 '\t\t(property "Reference" "R1"\n\t\t\t(at 0 0)\n\t\t)\n'
                 '\t\t(property "Value" "10k"\n\t\t\t(at 0 0)\n\t\t)\n'
                 '\t\t(property "MPN" "R0603-10K-TEST"\n\t\t\t(at 0 0)\n\t\t)\n'
                 '\t\t(attr smd dnp exclude_from_bom)\n\t)\n'
                 '\t(footprint "Logo:Logo"\n\t\t(layer "F.Cu")\n'
                 '\t\t(property "Reference" "G***"\n\t\t\t(at 0 0)\n\t\t)\n\t\t(attr board_only)\n\t)\n)\n')
        p = os.path.join(self.tmp, 'b.kicad_pcb')
        with open(p, 'w', encoding='utf-8', newline='\n') as f:
            f.write(board)
        fps = mx.board_footprints(p)
        self.assertEqual(fps[0], {'ref': 'R1', 'lib': 'Resistor_SMD:R_0603', 'layer': 'B.Cu', 'value': '10k',
                                  'mpn': 'R0603-10K-TEST', 'attrs': {'smd', 'dnp', 'exclude_from_bom'}})
        self.assertEqual(fps[1]['attrs'], {'board_only'})

    def test_netlist_comps(self):
        xml = ('<export version="E"><components>'
               '<comp ref="R1"><value>10k</value><footprint>R:R_0603</footprint>'
               '<fields><field name="MPN">RC</field></fields><property name="dnp"/></comp>'
               '<comp ref="R1"><value>10k</value></comp>'
               '<comp ref="TP1"><value>TP</value><property name="exclude_from_bom"/>'
               '<property name="exclude_from_board"/></comp>'
               '</components></export>')
        p = os.path.join(self.tmp, 'n.xml')
        with open(p, 'w', encoding='utf-8') as f:
            f.write(xml)
        comps, dups = mx.netlist_comps(p)
        self.assertEqual(dups, ['R1'])
        self.assertTrue(comps['TP1']['excl_bom'] and comps['TP1']['excl_board'])
        self.assertEqual(comps['R1']['value'], '10k')

    def test_project_revision(self):
        with open(os.path.join(self.tmp, 'b.kicad_pro'), 'w', encoding='utf-8') as f:
            f.write('{"text_variables": {"REVISION": "1.2"}}')
        self.assertEqual(mx.project_revision(self.tmp, 'b', None), '1.2')
        with open(os.path.join(self.tmp, 'c.kicad_pcb'), 'w', encoding='utf-8') as f:
            f.write('(kicad_pcb\n\t(property "REVISION" "3.4")\n)\n')
        self.assertEqual(mx.project_revision(self.tmp, 'c', os.path.join(self.tmp, 'c.kicad_pcb')), '3.4')


if __name__ == '__main__':
    unittest.main()
